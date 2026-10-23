"""Build shared notebooks, DOCX/PDF handouts and the maintained JavaScript slides.

Set SOFFICE to the prepared LibreOffice executable for PDF export.
Use --handouts-only when the slide deck is already built.
"""
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
ART = DOCS / "assets"
NAVY = "17354B"
TEAL = "197F80"
GREY = "607383"
PALE = "E4EFF0"
WHITE = "FFFFFF"
BG = "F7FAFB"
GOLD = "BF8D42"
RED = "DC6658"


def make_docx(filename, source, *, worksheet=False):
    """Render the Markdown handout source without maintaining a second lesson plan."""
    import re
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.top_margin = sec.bottom_margin = Inches(.45 if worksheet else .65)
    sec.left_margin = sec.right_margin = Inches(.58 if worksheet else .72)
    for name, size in (("Normal", 13 if filename == "activity-cards.docx" else 11), ("Title", 20 if worksheet else 24),
                       ("Heading 1", 11 if worksheet else 15), ("Heading 2", 10 if worksheet else 12)):
        st = d.styles[name]
        st.font.name, st.font.size = "DejaVu Sans", Pt(size)
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_after = Pt(3 if worksheet else 6)
    for style in d.styles:
        for border in style.element.findall(".//" + qn("w:pBdr")):
            border.getparent().remove(border)
    d.styles["Normal"].paragraph_format.line_spacing = 1.0 if worksheet else 1.08
    d.styles["Heading 1"].paragraph_format.space_before = Pt(7)
    d.styles["Title"].paragraph_format.space_after = Pt(6)
    d.styles["Title"].font.color.rgb = RGBColor(0, 0, 0)

    def clean(text):
        text = re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", text)
        return text.replace("**", "").replace("`", "")

    def runs(p, text, code=False):
        if code:
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(text)
            run.font.name, run.font.size = "Noto Sans Mono", Pt(9.5)
            return
        for i, part in enumerate(re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)):
            run = p.add_run(clean(part))
            if part.startswith("**"):
                run.bold = True
            if part.startswith("`"):
                run.font.name, run.font.size = "Noto Sans Mono", Pt(9.5)

    lines = (DOCS / source).read_text().splitlines()
    table = None
    code = False
    mission_titles = {m["title"] for m in __import__("build_notebooks").load_curriculum()["missions"]}
    mission_index = 0
    for line in lines:
        if line.startswith("<!--"):
            continue
        if line.startswith("```"):
            code = not code
            continue
        if not line.strip():
            table = None
            continue
        if code:
            runs(d.add_paragraph(), line, code=True)
            continue
        if line.startswith("!["):
            match = re.match(r"!\[([^]]*)\]\(([^)]+)\)", line)
            if match:
                d.add_picture(str((DOCS / match.group(2)).resolve()), width=Inches(6.6))
                continue
        if line.startswith("|"):
            values = [clean(v.strip()) for v in line.strip("|").split("|")]
            if all(re.fullmatch(r":?-+:?", v) for v in values):
                continue
            if table is None:
                table = d.add_table(rows=0, cols=len(values))
                table.style = "Light Shading Accent 1"
            cells = table.add_row().cells
            for cell, value in zip(cells, values):
                runs(cell.paragraphs[0], value)
                for run in cell.paragraphs[0].runs:
                    run.font.size = Pt(8)
            if len(table.rows) == 1:
                table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            continue
        if line.startswith("# "):
            d.add_paragraph(clean(line[2:]), "Title")
        elif line.startswith("### "):
            d.add_paragraph(clean(line[4:]), "Heading 2")
        elif line.startswith("## "):
            if line[3:] == "Optional color destinations extension" and filename == "student-handbook.docx":
                d.add_page_break()
            if line[3:] in mission_titles:
                if filename in ("activity-cards.docx", "instructor-guide.docx", "student-handbook.docx") or (worksheet and mission_index % 2 == 0):
                    d.add_page_break()
                mission_index += 1
            if filename == "student-handbook.docx" and line.startswith(("## A ", "## B ", "## C ", "## D ", "## E ", "## A ·", "## B ·", "## C ·", "## D ·", "## E ·")):
                d.add_page_break()
            d.add_paragraph(clean(line[3:]), "Heading 1")
        elif line.startswith("- [ ] "):
            runs(d.add_paragraph(), "[ ] " + line[6:])
        else:
            p = d.add_paragraph()
            runs(p, line, code=code)
    d.core_properties.title = next(line[2:] for line in lines if line.startswith("# "))
    d.core_properties.subject = "22 October 2026 workshop"
    footer = sec.footer.paragraphs[0]
    footer.alignment = 2
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    d.save(DOCS / filename)


def make_slides():
    """Delegate to the maintained JavaScript slide generator."""
    import subprocess
    subprocess.run(["node", str(ROOT / "tools/build_lecture_slides.mjs")], cwd=ROOT, check=True)
    return DOCS / "workshop-slides.pptx"


def export_handout_pdfs():
    """Use an explicit SOFFICE when the classroom build uses portable LibreOffice."""
    import os
    import shutil
    import subprocess
    import tempfile
    soffice = os.environ.get("SOFFICE") or shutil.which("soffice")
    if not soffice:
        raise RuntimeError("Set SOFFICE to the prepared LibreOffice executable to export handout PDFs")
    for name in ("instructor-guide", "student-worksheet", "student-handbook", "lesson-plan", "activity-cards"):
        with tempfile.TemporaryDirectory(prefix="workshop-handout-") as profile:
            subprocess.run([soffice, "-env:UserInstallation=" + Path(profile).as_uri(), "--headless", "--convert-to", "pdf", "--outdir", str(DOCS), str(DOCS/(name + ".docx"))], check=True, capture_output=True)
        if not (DOCS/(name + ".pdf")).is_file():
            raise RuntimeError("PDF export missing: " + name)


def normalize_artifact(path):
    """Remove exporter timestamps/IDs so repeated builds have identical bytes."""
    import re
    from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED
    import io
    if path.suffix in (".docx", ".pptx"):
        buffer = io.BytesIO()
        with ZipFile(path) as source, ZipFile(buffer, "w", ZIP_DEFLATED) as target:
            for name in sorted(source.namelist()):
                data = source.read(name)
                if name == "docProps/core.xml":
                    data = re.sub(rb"20[0-9]{2}-[0-9-]{5}T[0-9:]{8}Z", b"2026-10-22T00:00:00Z", data)
                info = ZipInfo(name, (2026,10,22,0,0,0))
                info.compress_type = ZIP_DEFLATED
                target.writestr(info, data)
        path.write_bytes(buffer.getvalue())
    elif path.suffix == ".pdf":
        data = path.read_bytes()
        # Same-length replacement preserves the PDF xref byte offsets.
        data = re.sub(rb"D:[0-9]{14}", b"D:20261022000000", data)
        data = re.sub(rb"20[0-9]{2}-[0-9-]{5}T[0-9:]{8}[+-][0-9:]{5}", b"2026-10-22T00:00:00+00:00", data)
        data = re.sub(rb"/DocChecksum /[A-Fa-f0-9]{32}", lambda m: b" " * len(m[0]), data)
        data = re.sub(rb"/ID\s*\[\s*<([A-Fa-f0-9]+)>\s*<([A-Fa-f0-9]+)>\s*\]",
                      lambda m: re.sub(rb"(?<=<)[A-Fa-f0-9]+(?=>)", lambda n: b"0"*len(n[0]), m[0]), data)
        path.write_bytes(data)


def bundle_receipt_inputs():
    import runpy
    builder=runpy.run_path(str(ROOT/'exercises/make_student_bundle.py'))
    return builder, builder['bundle_input_names'](ROOT)+['exercises/make_student_bundle.py','docs/student-workshop.zip']


def publish_bundle():
    """Build the tracked learner deliverable and hash every allowlisted input."""
    import hashlib,json
    builder,names=bundle_receipt_inputs()
    (DOCS/'student-workshop.zip').write_bytes(builder['bundle_bytes'](ROOT))
    receipt=json.loads((DOCS/'materials-build.json').read_text()) if (DOCS/'materials-build.json').exists() else {}
    names += ['tools/build_materials.py']
    receipt.update({name:hashlib.sha256((ROOT/name).read_bytes()).hexdigest() for name in names})
    (DOCS/'materials-build.json').write_text(json.dumps(receipt,indent=2,sort_keys=True)+'\n')


def main():
    import argparse
    from build_notebooks import main as notebooks
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--handouts-only", action="store_true")
    parser.add_argument("--check", action="store_true", help="Verify source projections and binary receipt without writing")
    parser.add_argument("--bundle-only",action="store_true",help="Rebuild published learner ZIP and receipt without rerendering materials")
    args = parser.parse_args()
    if args.bundle_only:
        notebooks(["--check"])
        publish_bundle()
        print("Published docs/student-workshop.zip and hashed all learner inputs")
        return
    if args.check:
        notebooks(["--check"])
        import hashlib, json
        receipt = json.loads((DOCS/"materials-build.json").read_text())
        builder,bundle_inputs=bundle_receipt_inputs()
        missing=set(bundle_inputs)-set(receipt)
        if missing: raise SystemExit("Untracked learner receipt inputs: "+", ".join(sorted(missing)))
        if (DOCS/'student-workshop.zip').read_bytes()!=builder['bundle_bytes'](ROOT):
            raise SystemExit("Learner bundle drift: docs/student-workshop.zip")
        for name, expected in receipt.items():
            if hashlib.sha256((ROOT/name).read_bytes()).hexdigest() != expected:
                raise SystemExit("Material drift: " + name)
        print("Material sources and binary hashes verified")
        return
    notebooks([])
    make_docx("instructor-guide.docx", "instructor.md")
    make_docx("student-worksheet.docx", "student-worksheet.md", worksheet=True)
    make_docx("student-handbook.docx", "student-exercises.md")
    make_docx("lesson-plan.docx", "lesson-plan.md")
    make_docx("activity-cards.docx", "activity-cards.md")
    export_handout_pdfs()
    if not args.handouts_only:
        make_slides()
    import hashlib, json
    names = ["curriculum/workshop.json", "tools/build_materials.py", "tools/build_lecture_slides.mjs", "tools/build_notebooks.py"]
    names += [str(p.relative_to(ROOT)) for p in sorted((DOCS/"assets/lecture").glob("*.png"))]
    names += ["docs/" + stem + ext for stem in ("instructor-guide", "student-worksheet", "student-handbook", "lesson-plan", "activity-cards", "workshop-slides") for ext in ((".pptx", ".pdf") if stem=="workshop-slides" else (".docx", ".pdf"))]
    for name in names:
        path = ROOT/name
        if path.suffix in (".docx", ".pdf", ".pptx") and path.exists(): normalize_artifact(path)
    (DOCS/"materials-build.json").write_text(json.dumps({name:hashlib.sha256((ROOT/name).read_bytes()).hexdigest() for name in names if (ROOT/name).exists()}, indent=2)+"\n")
    publish_bundle()
    print("Built handouts, notebooks and published learner ZIP in", DOCS, "and", ROOT/"notebooks")


if __name__ == "__main__":
    main()
